import os
import random
import openpyxl
import json
import yaml
import encryption

# --- Configuration and Paths Setup ---
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(CURRENT_DIR)

yaml_path = os.path.join(CURRENT_DIR, 'config.yaml')
with open(yaml_path, 'r', encoding='utf-8') as file:
    config = yaml.safe_load(file)

INPUT_DIR = os.path.join(BASE_DIR, "InputFiles")
OUTPUT_DIR = os.path.join(BASE_DIR, "OutputFiles")

os.makedirs(OUTPUT_DIR, exist_ok=True)


def load_file_lines(filepath):
    """Loads lines from a text file, ignoring blank spaces."""
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            content = f.read()
            return [line.strip() for line in content.splitlines() if line.strip()]
    except FileNotFoundError:
        print(f"Error: The file '{filepath}' was not found.")
        return []

def create_synoptic_table(medications):
    """Generates an interaction matrix for a list of medications."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Synoptic Table"

    interaction_rules = {}

    for i, name_i in enumerate(medications, start=2):
        ws.cell(row=i, column=1, value=name_i)
        ws.cell(row=1, column=i, value=name_i)

        for j, name_j in enumerate(medications, start=2):
            if i == j:
                value = 0
            else:
                value = random.randint(0, 6)

            ws.cell(row=i, column=j, value=value)
            interaction_rules[(name_i, name_j)] = value

    return wb, interaction_rules

def save_synoptic_table(filename, wb):
    """Saves the synoptic table to Excel and Tab-separated values."""
    excel_path = os.path.join(OUTPUT_DIR, f"{filename}.xlsx")
    txt_path = os.path.join(OUTPUT_DIR, f"{filename}.txt")
    
    try:
        wb.save(excel_path)

        with open(txt_path, "w", encoding="utf-8") as f:
            for row in wb.active.iter_rows(values_only=True):
                f.write("\t".join(str(cell) if cell is not None else "" for cell in row) + "\n")

        print(f"Synoptic table successfully saved as '{filename}.xlsx' and '{filename}.txt'.")

    except Exception as e:
        print(f"Error saving synoptic table: {e}")

def generate_final_data(names_path, surnames_path, num_patients, medications_list, rules, key):
    """Generates patients profiles with prescriptions and evaluates drug interactions."""
    base_names = load_file_lines(names_path)
    base_surnames = load_file_lines(surnames_path)

    id_list = []
    json_list = []

    while len(id_list) < num_patients:
        new_id = random.randint(1000, 9999)
        if new_id not in id_list:
            id_list.append(new_id)

    for uid in id_list:
        full_name = f"{random.choice(base_names)} {random.choice(base_surnames)}"
        prescription = random.sample(medications_list, k=random.randint(3, 5))

        total_interactions = 0
        has_interaction = False

        for idx in range(len(prescription)):
            for j in range(idx + 1, len(prescription)):
                severity = rules.get((prescription[idx], prescription[j]), 0)
                total_interactions += severity
                if severity > 0:
                    has_interaction = True

        encrypted_id = encryption.encrypt(str(uid), key)

        json_list.append({
            "patient_id": encrypted_id,
            "patient_name": full_name,
            "prescription": prescription,
            "total_balance": total_interactions,
            "result": "With Interactions" if has_interaction else "No Interactions"
        })

    return json_list

def save_json(data, filename):
    """Saves the final generated patient profiles to a JSON file."""
    filepath = os.path.join(OUTPUT_DIR, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4, ensure_ascii=False)

    print(f"JSON profile data successfully created at '{filename}'.")

def main():
    """Main execution pipeline."""
    num_patients = config['parameters']['num_patients']
    encryption_key = config['parameters']['default_key']
    
    meds_file = config['input_files']['medications']
    names_file = config['input_files']['first_names']
    surnames_file = config['input_files']['last_names']
    
    medications_list = load_file_lines(os.path.join(INPUT_DIR, meds_file))
    if not medications_list:
        return
        
    wb, rules = create_synoptic_table(medications_list)
    
    excel_base_name = config['output_files']['synoptic_excel'].replace('.xlsx', '')
    save_synoptic_table(excel_base_name, wb)

    names_path = os.path.join(INPUT_DIR, names_file)
    surnames_path = os.path.join(INPUT_DIR, surnames_file)
    
    final_data = generate_final_data(
        names_path, surnames_path, num_patients, medications_list, rules, encryption_key
    )
    
    json_filename = config['output_files']['final_json']
    save_json(final_data, json_filename)

if __name__ == "__main__":  
    main()